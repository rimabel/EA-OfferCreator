#!/usr/bin/env python3
"""
Equipment: reise_pdf_erstellen_linux.py
Zweck:  Kombiniertes Reise-PDF aus sights.json + tagesplan.json erstellen.
        Seite 1: Firmen-Header + Tagesablauf als Stop-Cards mit Bullet Points
        Seite 2+: Sehenswürdigkeiten mit Fotos
        Ausgabe: BR-REISE-[ZIELORT]_Garamond_4_.docx + BR-REISE-[ZIELORT].pdf
        PDF-Konvertierung: LibreOffice (plattformunabhängig, läuft auf Linux)
Voraussetzung:
  sudo apt install libreoffice  # oder: dnf install libreoffice
Aufruf:
  python equipment/reise_pdf_erstellen_linux.py \\
      --ziel "Paris" \\
      --sights "equipment/temp/reise-paris/sights.json" \\
      --tagesplan "equipment/temp/reise-paris/tagesplan.json"
"""

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    from PIL import Image as PilImage
    import io as _io
    PILLOW_AVAILABLE = True
except ImportError:
    PILLOW_AVAILABLE = False

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

COLOR_DARK_RED  = RGBColor(0x98, 0x00, 0x00)
COLOR_BURGUNDY  = RGBColor(0x6E, 0x14, 0x14)
COLOR_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GRAY_TEXT = RGBColor(0x66, 0x66, 0x66)
COLOR_DARK_GRAY = RGBColor(0x33, 0x33, 0x33)

COL_HEADER_LOGO_W = Cm(3)
COL_IMAGE_W       = Cm(14)
MAX_IMG_HEIGHT    = Cm(23)

COMPANY_NAME = "BELAHMER REISEN"
COMPANY_LINES = [
    "Inh. Nabil Belahmer",
    "Kurpfalzstr 3",
    "67734 Katzweiler",
    "E-Mail: info@belahmer-reisen.de",
    "Tel.: +49 6301 6689790",
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def clean_description(text: str) -> str:
    if not text:
        return text
    text = re.sub(
        r'\s*\([^)]*(?:französisch|deutsch|englisch|lateinisch|arabisch|spanisch|italienisch)[^)]*\)',
        '', text
    )
    text = re.sub(r'\s*\[[^\]]*\]', '', text)
    text = re.sub(r'…$', '', text)
    text = re.sub(r',\s*,', ',', text)
    text = re.sub(r'\s{2,}', ' ', text).strip().rstrip(',').strip()
    return text


def set_keep_with_next(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement('w:keepNext')
    pPr.append(keep)


def normalize_ziel(ziel: str) -> str:
    decomposed = unicodedata.normalize("NFKD", ziel)
    no_accents = "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")
    umlaut_table = str.maketrans({
        "ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss",
        "Ä": "ae", "Ö": "oe", "Ü": "ue",
    })
    normalized = no_accents.translate(umlaut_table)
    return normalized.lower().replace(" ", "-")


def filename_ziel(ziel: str) -> str:
    decomposed = unicodedata.normalize("NFKD", ziel)
    no_accents = "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")
    umlaut_table = str.maketrans({
        "ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss",
        "Ä": "Ae", "Ö": "Oe", "Ü": "Ue",
    })
    normalized = no_accents.translate(umlaut_table)
    return normalized.replace(" ", "-")


def set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    existing = tcPr.find(qn("w:shd"))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(shd)


def set_cell_no_border(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        tcBorders.append(border)
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcBorders)


def set_table_no_border(table) -> None:
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        tblBorders.append(border)
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(tblBorders)


def set_paragraph_space(paragraph, before_pt: int = 0, after_pt: int = 0) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    if before_pt:
        spacing.set(qn("w:before"), str(before_pt * 20))
    if after_pt:
        spacing.set(qn("w:after"), str(after_pt * 20))


def add_divider(doc: Document) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "888888")
    pBdr.append(bottom)
    pPr.append(pBdr)
    set_paragraph_space(p, before_pt=4, after_pt=4)
    set_keep_with_next(p)


def add_stop_heading(doc: Document, text: str, italic: bool = False) -> None:
    p = doc.add_paragraph()
    set_paragraph_space(p, before_pt=4, after_pt=2)
    run = p.add_run(text)
    run.bold = True
    run.italic = italic
    run.font.size = Pt(12)
    run.font.color.rgb = COLOR_BURGUNDY


def add_time_range(doc: Document, von: str, bis: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_space(p, before_pt=0, after_pt=2)
    run = p.add_run(f"{von} – {bis}")
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_GRAY_TEXT


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_space(p, before_pt=2, after_pt=2)
    pPr = p._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "227")
    pPr.append(ind)
    bullet_run = p.add_run("• " + text)
    bullet_run.font.size = Pt(10)
    bullet_run.font.color.rgb = COLOR_DARK_GRAY


def add_sight_description(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_space(p, before_pt=2, after_pt=6)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)


def add_return_line(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_space(p, before_pt=2, after_pt=2)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_GRAY_TEXT


def add_watermark_to_header(section, color: str = "D4A8A8", text: str = "BELAHMER REISEN") -> None:
    from lxml import etree
    header = section.header
    header.is_linked_to_previous = False
    for para in header.paragraphs:
        para.clear()
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    watermark_xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        '<w:rPr><w:noProof/></w:rPr>'
        '<w:pict>'
        f'<v:shape id="WaterMark" type="#_x0000_t136"'
        f' style="position:absolute;margin-left:0;margin-top:0;'
        f'width:527.85pt;height:65.95pt;z-index:-251656192;'
        f'mso-wrap-style:none;'
        f'mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        f'mso-position-vertical:center;mso-position-vertical-relative:margin;'
        f'rotation:315"'
        f' fillcolor="#{color}" stroked="f">'
        f'<v:fill on="t" focussize="0,0"/>'
        f'<v:path textpathok="t"/>'
        f'<v:textpath on="t" string="{text}"'
        f' style="font-family:&quot;Calibri&quot;;font-size:1pt"/>'
        f'</v:shape>'
        '</w:pict>'
        '</w:r>'
    )
    para._p.append(etree.fromstring(watermark_xml))


def add_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    for para in footer.paragraphs:
        para.clear()
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top_border = OxmlElement("w:top")
    top_border.set(qn("w:val"), "single")
    top_border.set(qn("w:sz"), "6")
    top_border.set(qn("w:space"), "4")
    top_border.set(qn("w:color"), "6E1414")
    pBdr.append(top_border)
    pPr.append(pBdr)
    run = para.add_run(
        "Belahmer Reisen | Inh. Nabil Belahmer | Kurpfalzstr 3 | "
        "67734 Katzweiler | info@belahmer-reisen.de | +49 6301 6689790 | Seite "
    )
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run2 = para.add_run()
    run2.font.name = "Calibri"
    run2.font.size = Pt(8)
    r_elem = run2._r
    fldChar_begin = OxmlElement("w:fldChar")
    fldChar_begin.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.text = "PAGE"
    fldChar_sep = OxmlElement("w:fldChar")
    fldChar_sep.set(qn("w:fldCharType"), "separate")
    fldChar_end = OxmlElement("w:fldChar")
    fldChar_end.set(qn("w:fldCharType"), "end")
    r_elem.append(fldChar_begin)
    r_elem.append(instrText)
    r_elem.append(fldChar_sep)
    r_elem.append(fldChar_end)


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------

def build_document(ziel: str, sights: list, tagesplan: dict, logo_path: str, datum: str = "", watermark: bool = False) -> Document:
    doc = Document()

    section = doc.sections[0]
    section.page_width  = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin    = Cm(1)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(2)
    section.right_margin  = Cm(2)

    header_table = doc.add_table(rows=1, cols=2)
    set_table_no_border(header_table)
    header_table.autofit = False
    header_table.columns[0].width = COL_HEADER_LOGO_W
    header_table.columns[1].width = Cm(14)

    logo_cell = header_table.cell(0, 0)
    info_cell = header_table.cell(0, 1)
    set_cell_no_border(logo_cell)
    set_cell_no_border(info_cell)

    for cell in (logo_cell, info_cell):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        vAlign = OxmlElement("w:vAlign")
        vAlign.set(qn("w:val"), "center")
        existing = tcPr.find(qn("w:vAlign"))
        if existing is not None:
            tcPr.remove(existing)
        tcPr.append(vAlign)

    logo_para = logo_cell.paragraphs[0]
    logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_run = logo_para.add_run()
    logo_run.add_picture(logo_path, width=Cm(3.2))

    info_para = info_cell.paragraphs[0]
    info_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    name_run = info_para.add_run(COMPANY_NAME)
    name_run.bold = True
    name_run.font.size = Pt(14)
    name_run.font.color.rgb = COLOR_DARK_RED

    for line in COMPANY_LINES:
        info_para.add_run().add_break()
        line_run = info_para.add_run(line)
        line_run.font.size = Pt(10)
        line_run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    title_para = doc.add_paragraph()
    set_paragraph_space(title_para, before_pt=24, after_pt=6)
    title_text = f"Tagesablauf: {ziel}  {datum}".strip() if datum else f"Tagesablauf: {ziel}"
    title_run = title_para.add_run(title_text)
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = COLOR_BURGUNDY

    ankunft = tagesplan.get("ankunft", {})
    ankunft_zeit = ankunft.get("zeit", "")
    ankunft_aktivitaeten = ankunft.get("aktivitaeten", [])

    add_divider(doc)
    add_stop_heading(doc, f"{ankunft_zeit} – Ankunft in {ziel}")
    for aktivitaet in ankunft_aktivitaeten:
        add_bullet(doc, aktivitaet)

    halte = tagesplan.get("halte", [])
    for halt in halte:
        nr = halt.get("nr", "")
        name = halt.get("name", "")
        von = halt.get("von", "")
        bis = halt.get("bis", "")
        aktivitaeten = halt.get("aktivitaeten", [])
        add_divider(doc)
        add_stop_heading(doc, f"{nr}. Halt: {name}")
        add_time_range(doc, von, bis)
        for aktivitaet in aktivitaeten:
            add_bullet(doc, aktivitaet)

    optional = tagesplan.get("optional")
    if optional:
        opt_name = optional.get("name", "")
        opt_von = optional.get("von", "")
        opt_bis = optional.get("bis", "")
        opt_aktivitaeten = optional.get("aktivitaeten", [])
        add_divider(doc)
        add_stop_heading(doc, f"Optionaler Halt: {opt_name}", italic=True)
        add_time_range(doc, opt_von, opt_bis)
        for aktivitaet in opt_aktivitaeten:
            add_bullet(doc, aktivitaet)

    abfahrt = tagesplan.get("abfahrt", {})
    abfahrt_zeit = abfahrt.get("zeit", "")
    abfahrt_heimatort = abfahrt.get("heimatort", "")
    heimankunft = tagesplan.get("heimankunft")

    add_divider(doc)
    if abfahrt_zeit and abfahrt_heimatort:
        add_return_line(doc, f"{abfahrt_zeit} – Rückfahrt nach {abfahrt_heimatort}")
    if heimankunft:
        heim_zeit = heimankunft.get("zeit", "")
        heim_ort = heimankunft.get("heimatort", "")
        if heim_zeit and heim_ort:
            add_return_line(doc, f"{heim_zeit} – Geplante Ankunft in {heim_ort}")

    from docx.enum.section import WD_SECTION_START
    doc.add_section(WD_SECTION_START.NEW_PAGE)

    section1 = doc.sections[0]
    section2 = doc.sections[1]

    section2.page_width    = Cm(21)
    section2.page_height   = Cm(29.7)
    section2.top_margin    = Cm(1)
    section2.bottom_margin = Cm(1.5)
    section2.left_margin   = Cm(2)
    section2.right_margin  = Cm(2)

    if watermark:
        add_watermark_to_header(section1, color="EDD8D8", text="BELAHMER REISEN")

    section2.header.is_linked_to_previous = False
    add_footer(section1)
    section2.footer.is_linked_to_previous = True

    section_title = doc.add_paragraph()
    set_paragraph_space(section_title, before_pt=0, after_pt=6)
    st_run = section_title.add_run("Sehenswürdigkeiten")
    st_run.bold = True
    st_run.font.size = Pt(14)
    st_run.font.color.rgb = COLOR_BURGUNDY

    for sight in sights:
        name = sight.get("name", "")
        image_path = sight.get("image")
        description = clean_description(sight.get("description"))

        sight_tbl = doc.add_table(rows=1, cols=1)
        set_table_no_border(sight_tbl)
        set_cell_no_border(sight_tbl.rows[0].cells[0])
        trPr = sight_tbl.rows[0]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trPr.append(cant_split)
        cell = sight_tbl.rows[0].cells[0]

        name_para = cell.paragraphs[0]
        set_paragraph_space(name_para, before_pt=8, after_pt=4)
        name_run = name_para.add_run(name)
        name_run.bold = True
        name_run.font.size = Pt(12)
        name_run.font.color.rgb = COLOR_BURGUNDY

        image_ok = (
            image_path is not None
            and isinstance(image_path, str)
            and os.path.isfile(image_path)
        )

        if image_ok:
            img_para = cell.add_paragraph()
            img_run = img_para.add_run()
            if PILLOW_AVAILABLE:
                try:
                    with PilImage.open(image_path) as pil_img:
                        img_w_px, img_h_px = pil_img.size
                        buf = _io.BytesIO()
                        pil_img.convert("RGB").save(buf, format="JPEG", quality=90)
                        buf.seek(0)
                    aspect = img_w_px / img_h_px if img_h_px else 1
                    display_h_cm = 14 / aspect
                    if display_h_cm > 23:
                        img_run.add_picture(buf, height=MAX_IMG_HEIGHT)
                    else:
                        img_run.add_picture(buf, width=COL_IMAGE_W)
                except (OSError, ValueError, TypeError) as pil_err:
                    print(f"  [WARN] Pillow-Konvertierung fehlgeschlagen ({pil_err}), fallback wird verwendet...")
                    img_run.add_picture(image_path, width=COL_IMAGE_W)
            else:
                img_run.add_picture(image_path, width=COL_IMAGE_W)
            set_paragraph_space(img_para, after_pt=2)
        else:
            placeholder = cell.add_paragraph()
            set_paragraph_space(placeholder, after_pt=2)
            ph_run = placeholder.add_run("[Kein Bild verfügbar]")
            ph_run.italic = True
            ph_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

        if description:
            desc_para = cell.add_paragraph()
            set_paragraph_space(desc_para, before_pt=2, after_pt=6)
            desc_run = desc_para.add_run(description)
            desc_run.italic = True
            desc_run.font.size = Pt(10)
            desc_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    return doc


# ---------------------------------------------------------------------------
# Argument parser
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Reise-PDF aus sights.json + tagesplan.json erstellen (Linux/LibreOffice)."
    )
    parser.add_argument("--ziel", required=True,
                        help="Reiseziel Anzeigename (z.B. 'Paris')")
    parser.add_argument("--sights", required=True,
                        help="Pfad zur sights.json")
    parser.add_argument("--tagesplan", required=True,
                        help="Pfad zur tagesplan.json")
    parser.add_argument("--datum", default="",
                        help="Reisetag (z.B. '22.05.2026')")
    parser.add_argument("--watermark", action="store_true",
                        help="Wasserzeichen 'BELAHMER REISEN' auf Seite 1 einblenden")
    return parser.parse_args()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    args = parse_args()

    if not os.path.isfile(args.sights):
        print(f"[ERROR] sights.json nicht gefunden: {args.sights}")
        sys.exit(1)

    if not os.path.isfile(args.tagesplan):
        print(f"[ERROR] tagesplan.json nicht gefunden: {args.tagesplan}")
        sys.exit(1)

    try:
        with open(args.sights, "r", encoding="utf-8") as f:
            sights = json.load(f)
    except json.JSONDecodeError as e:
        print(f"[ERROR] sights.json hat kein gueltiges JSON-Format: {e}")
        sys.exit(1)

    try:
        with open(args.tagesplan, "r", encoding="utf-8") as f:
            tagesplan = json.load(f)
    except json.JSONDecodeError as e:
        print(f"[ERROR] tagesplan.json hat kein gueltiges JSON-Format: {e}")
        sys.exit(1)

    logo_path1 = os.path.join(SCRIPT_DIR, "template", "word", "media", "image1.png")
    logo_path = logo_path1
    if not os.path.isfile(logo_path):
        project_root = os.path.dirname(SCRIPT_DIR)
        logo_path2 = os.path.join(project_root, "template", "word", "media", "image1.png")
        print(f"  [INFO] Logo nicht gefunden unter {logo_path1}, versuche {logo_path2}...")
        logo_path = logo_path2

    if not os.path.isfile(logo_path):
        print(f"[ERROR] Logo nicht gefunden: {logo_path}")
        sys.exit(1)

    ziel_for_file = filename_ziel(args.ziel)
    cwd = os.getcwd()
    docx_filename = f"BR-REISE-{ziel_for_file}_Garamond_4_.docx"
    pdf_filename  = f"BR-REISE-{ziel_for_file}.pdf"

    docx_path = os.path.join(cwd, docx_filename)
    pdf_path  = os.path.join(cwd, pdf_filename)

    # --- Build DOCX ---
    doc = build_document(
        ziel=args.ziel,
        sights=sights,
        tagesplan=tagesplan,
        logo_path=logo_path,
        datum=args.datum,
        watermark=args.watermark,
    )
    doc.save(docx_path)
    print(f"[OK] DOCX erstellt: {docx_filename}")

    # --- Convert to PDF via LibreOffice (headless) ---
    try:
        result = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf",
             "--outdir", cwd, docx_path],
            capture_output=True,
            text=True,
            timeout=120,
        )
        if result.returncode != 0:
            print(f"[ERROR] LibreOffice Konvertierung fehlgeschlagen:\n{result.stderr}")
            sys.exit(1)

        # LibreOffice benennt die Ausgabedatei nach dem Eingabedateinamen
        # BR-REISE-Paris_Garamond_4_.docx → BR-REISE-Paris_Garamond_4_.pdf
        # Wir benennen sie in BR-REISE-Paris.pdf um.
        lo_pdf = os.path.join(cwd, os.path.splitext(docx_filename)[0] + ".pdf")
        if os.path.isfile(lo_pdf) and lo_pdf != pdf_path:
            shutil.move(lo_pdf, pdf_path)

    except FileNotFoundError:
        print("[ERROR] LibreOffice nicht gefunden. Bitte installieren:")
        print("  Ubuntu/Debian: sudo apt install libreoffice")
        print("  Fedora/RHEL:   sudo dnf install libreoffice")
        sys.exit(1)
    except subprocess.TimeoutExpired:
        print("[ERROR] LibreOffice Konvertierung hat das Zeitlimit (120s) überschritten.")
        sys.exit(1)
    except Exception as e:
        print(f"[ERROR] Fehler bei der PDF-Konvertierung: {e}")
        sys.exit(1)

    print(f"[OK] PDF erstellt: {pdf_filename}")
